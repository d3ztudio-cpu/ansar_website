from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r'G:\ANSAR WEBSITE\ansar_website\ansar_website\Library_Admin_Panel_Instruction_Manual.docx'
doc = Document(); s=doc.sections[0]; s.top_margin=Inches(.7); s.bottom_margin=Inches(.7); s.left_margin=Inches(.8); s.right_margin=Inches(.8)
for n,sz,c in [('Title',28,'15324A'),('Heading 1',19,'15324A'),('Heading 2',14,'0F766E'),('Heading 3',11,'15324A')]:
    st=doc.styles[n]; st.font.name='Aptos Display'; st.font.size=Pt(sz); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(c); st._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos Display')
doc.styles['Normal'].font.name='Aptos'; doc.styles['Normal'].font.size=Pt(10)

def shade(cell,fill):
    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),fill); cell._tc.get_or_add_tcPr().append(sh)
def tbl(rows):
    t=doc.add_table(rows=0,cols=len(rows[0])); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for ri,row in enumerate(rows):
        cs=t.add_row().cells
        for i,v in enumerate(row):
            cs[i].text=str(v); cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cs[i].paragraphs:
                p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2)
                for r in p.runs: r.font.size=Pt(8.5)
            if ri==0:
                shade(cs[i],'15324A')
                for p in cs[i].paragraphs:
                    for r in p.runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
            elif ri%2==0: shade(cs[i],'F4F8FA')
    return t
def para(x='',style=None):
    p=doc.add_paragraph(style=style); p.add_run(x); p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.08; return p
def bullets(xs):
    for x in xs: para(x,'List Bullet')
def steps(xs):
    for x in xs: para(x,'List Number')
def page(): doc.add_page_break()

doc.add_paragraph('ANSAR ENGLISH SCHOOL',style='Subtitle').alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Library Admin Panel Instruction Manual',style='Title').alignment=WD_ALIGN_PARAGRAPH.CENTER
para('Complete operating guide for library content editors').alignment=WD_ALIGN_PARAGRAPH.CENTER
para('Version 1.0 | Prepared 31 October 2026').alignment=WD_ALIGN_PARAGRAPH.CENTER
para('Purpose: This manual explains every control in the Library section of the administrator panel: uploading, editing, publishing, unpublishing, deleting, moving, searching, bulk creation, profile management, and troubleshooting.')
tbl([['Quick fact','How the system works'],['Admin location','Admin panel -> Library.'],['Database','librarySettings, libraryPosts, and libraryOneFact in Firestore.'],['Images','School server upload service; maximum 15 MB per image.'],['PDFs','Authenticated PDF upload function; maximum 25 MB per PDF.'],['Public visibility','Only records with Published enabled appear publicly.'],['Deletion warning','Removing a record removes the database reference; physical uploaded files are immutable and remain on server storage.']])
page()

doc.add_heading('Contents and panel map',1)
tbl([['Tab','Purpose'],['Overview','Counts, drafts, section totals, and public-page shortcut.'],['Library Profile','Title, copy, hours, contact, rules, vision, mission, quiz URL, and staff JSON.'],['One Fact a Day','Date-based daily posters, archive import, and bulk poster creation.'],['New Arrivals','New books/publications, status, accession number, and cover.'],['Question Papers','Class, subject, category, year, and Drive document.'],['E-Books','Digital books, reader links, covers, and automatic AES-E numbers.'],['E-Learning','Learning platforms and resource links.'],['Short Films','YouTube, Vimeo, and other video links.'],['Creative Corner','Student/staff reviews, writing, art, and creative projects.'],['Quizzes','Quiz URLs and optional instructions.'],['Events & Activities','Library events and activities with date and venue.'],['Winners & Gallery','Winner entries with multiple photos.'],['Newspaper Clippings','Press coverage and newspaper entries.'],['Files & Documents','PDFs, circulars, reference documents, and external files.'],['Vocabulary Vault','One dated vocabulary poster per day.'],['Published','All library posts together, including drafts.']])
doc.add_heading('Access and roles',2)
para('Sign in to the school administrator area and select Library. The library-only account is restricted to this module. Full administrators may see other modules. Public visitors can read only published content; drafts and management controls are private.')
doc.add_heading('Common list controls',2)
bullets(['Edit fills the form with the selected record and scrolls to the form. Select Update item or Update entry to save.','Delete asks for confirmation and permanently deletes the Firestore record. There is no undo button.','Cancel editing discards unsaved form changes.','Published and Draft labels identify visibility.','The live Firestore listener refreshes lists after successful writes.'])

doc.add_heading('Publishing model',1)
tbl([['Action','Result'],['Publish checked','published=true; item appears on the public Library page.'],['Publish unchecked','published=false; item remains visible to authorised administrators only.'],['Unpublish an existing item','Edit it, clear Publish, and save.'],['Delete','Deletes the record after confirmation. It does not remove the physical upload.'],['Legacy Google Sites URL','Blocked at save; replace it with a current URL or upload.']])

doc.add_heading('Overview tab',1)
para('Overview is a monitoring dashboard, not a content form.')
bullets(['At a glance shows total Library items, draft count, question papers, and documents.','Items per section groups records by normalised type, including older records stored under legacy types.','Open public Library page opens /library in a new tab.'])
steps(['Wait for counts to load.','Compare totals with the expected work.','Open the public page after publishing to verify the exact card, section, image, and link.'])

doc.add_heading('Library Profile tab',1)
para('Select Save library profile after changing any field. The values are stored in the single librarySettings/main document.')
tbl([['Field','Instruction'],['Page title','Required main heading.'],['Subtitle','Required supporting line.'],['Introduction','Required introductory paragraph.'],['Hero image URL','Optional valid image URL.'],['Library hours','Complete display text for opening hours.'],['Library contact email','Optional valid email.'],['Library days','Operating days.'],['Digital quiz URL','Link for the library quiz entry point.'],['Vision / Mission','Public vision and mission text.'],['Library rules','One rule per line; lines are displayed separately.'],['Library staff directory JSON','Array of objects with name, role, qualification, contact, and optional imageUrl.'],['Save library profile','Writes the profile and updates the public page.']])
doc.add_heading('Staff JSON example and rules',2)
para('[{"name":"SHYNI PK","role":"Librarian","qualification":"MLIS","contact":"9400566641","imageUrl":""}]')
bullets(['Use a JSON array with square brackets. Use double quotes and valid commas.','Validate JSON before saving; malformed JSON can break staff display.','Do not place passwords or private student data in this field.'])

doc.add_heading('One Fact a Day tab',1)
para('Records are stored in libraryOneFact. The document ID is the poster date YYYY-MM-DD, so one date can contain only one entry.')
tbl([['Control','Instruction'],['Poster date','Required date represented by the poster.'],['Publish this poster','Controls public visibility.'],['Title / Description','Optional text.'],['Poster image','One image, maximum 15 MB.'],['Save / Update entry','Creates or updates the date record.']])
steps(['Choose a date. If it already exists, edit it rather than creating a duplicate.','Enter optional text and upload the image.','Choose Published or Draft.','Select Save entry and verify the archive list.'])
doc.add_heading('Archive import',2)
para('The panel can create dated entries for the 143 local posters (1.jpg to 143.jpg). The date entered for poster 143 is used as the archive endpoint; existing dates are skipped.')
steps(['Set Date for poster 143.','Select Import 143 posters.','Read the number imported. If every date exists, the panel reports that nothing was imported.'])
doc.add_heading('Bulk poster upload',2)
steps(['Select or drop multiple images.','Set the start date for the first poster.','Review the natural filename order.','Select Create entries for N posters.','The system moves forward by date and skips dates already used.'])
para('The archive list shows the 60 most recent entries. Older entries remain available by date in the edit form.')

doc.add_heading('Vocabulary Vault tab',1)
para('Vocabulary Vault is a libraryPosts record with type Vocabulary. It is separate from One Fact a Day.')
bullets(['Poster date is required and duplicate dates are prevented when creating.','Title normally contains the word; the image is uploaded through the image uploader.','The dedicated form saves the entry as published.','Edit changes date, title, or image. Delete removes the post record.'])

doc.add_heading('Shared form and media uploader',1)
tbl([['Control','Behavior'],['Item type','Selects the record type; Events use Event/Activity buttons.'],['Date','Optional on most tabs; required for Events and Activities.'],['Title','Required standard field.'],['Description','Optional summary where shown.'],['Image','Single optional cover/thumbnail unless a gallery is used.'],['Publish on the Library page','Public/private switch.'],['Save / Update item','Writes the record. Disabled while data is loading.']])
doc.add_heading('Image upload procedure',2)
steps(['Select or drop an image. In multi-upload mode, select several.','Wait for the progress indicator to finish.','Review thumbnails. Use Remove to remove a URL from the record or Replace to choose a new single image.','If an upload fails, select Retry upload.','Save the record after the correct URL appears in the form.'])
bullets(['Only images are accepted; each must be 15 MB or smaller.','The upload service validates the real file content even when a device reports a missing MIME type.','Removing an image from the form does not delete the physical server file.'])
doc.add_heading('PDF upload procedure',2)
steps(['Select a PDF, maximum 25 MB.','Remain signed in with an authorised account.','Wait for the authenticated upload to finish.','Confirm the saved URL or Drive sharing link.','Select Save item.'])
para('If the PDF uploader requests sign-in, sign in again before retrying. If the upload succeeded but the later Firestore save failed, retry the save rather than uploading duplicate copies.')

tabs=[
('New Arrivals','New Arrival / Publication','Title of book, author, date, book status NEW or EXISTING, optional accession number, cover image, Publish. NEW receives a red public cover stamp. Items can be moved to Events & Activities.'),
('Question Papers','Question Paper','Class, Subject, Exam category, Year, Title, and required Google Drive link. Set Drive access to Anyone with the link can view. Public filters use class, subject, category, and year.'),
('E-Books','E-Book','Title and required reader/viewer link. Optional description, date, accession number, and cover. Blank number generates the next AES-E number. The list can search title or book number.'),
('E-Learning','E-Learning','Title, required resource link, and optional description. Date and image are hidden. Links open in a new tab.'),
('Short Films','Short Film','Title and required video URL. YouTube/Vimeo embed; other approved URLs open externally. Date and image are hidden.'),
('Creative Corner','Creative Corner','Contributor type Student/Staff; Name; Section; student Class & division; Title; Date; optional image and video URL. Required values are enforced for the selected contributor.'),
('Quizzes','Quiz','Title and required quiz URL. Optional instructions/description, date, and cover image.'),
('Events & Activities','Event / Activity','Choose the type, then enter required Date, Title, Description, Venue, optional image, and Publish. The type is normalised for public placement.'),
('Winners & Gallery','Winner','Title, date, optional category, winner/newspaper field, venue/source, and multiple photos. The public page shows the photos as a collage.'),
('Newspaper Clippings','Clipping','Title, date, optional description, image, category/programme, author/newspaper, venue/source, and Publish.'),
('Files & Documents','File','File type, category/collection, required PDF upload or Google Drive link, title, optional date, cover thumbnail, and Publish. File types: PDF, Image, Google Drive document, Web link, Other.'),
]
for name,typ,fields in tabs:
    doc.add_heading(name+' tab',1); tbl([['Record type','Controls and field rules'],[typ,fields]])
    steps(['Open the tab and confirm the type.','Complete required fields and links.','Upload and review media if applicable.','Choose Publish or leave as draft.','Save and verify the list, then check the public page.'])

doc.add_heading('Published tab',1)
para('Published is an all-post review list. It includes published items and drafts across standard library post sections. It does not replace the One Fact a Day or Profile tabs.')
bullets(['Use it when the original section is unknown.','Rows show normalised type, status, class/year where available, title, date, and subject where available.','Edit or Delete directly from the row.','Drafts are intentionally shown so they can be reviewed and released.'])

doc.add_heading('Editing, deleting, and moving',1)
doc.add_heading('Edit',2)
steps(['Select Edit on the item.','Review every pre-filled field, especially type, date, link, and publication status.','Replace retired links and correct metadata.','Select Update item or Update entry.','Open the public page and verify the result.'])
doc.add_heading('Delete',2)
steps(['Select Delete.','Read the confirmation prompt.','Confirm only when the record should permanently disappear.','Verify removal from the list and public page.'])
para('Deletion cannot be undone in the panel. It removes the Firestore record, not the physical uploaded asset.')
doc.add_heading('Move New Arrival to Events & Activities',2)
para('Select Move to Events & Activities and confirm. The record changes to Event and its book status becomes EXISTING. Edit afterward if it should be Activity or needs a different date/venue.')

doc.add_heading('Validation and troubleshooting',1)
tbl([['Message / symptom','Action'],['Only image files are allowed','Choose a real image file.'],['Only PDF files can be uploaded','Choose a PDF.'],['Image over 15 MB / PDF over 25 MB','Compress or resize before retrying.'],['Please sign in again before uploading a PDF','Sign out/in with an authorised account.'],['Paste the Google Drive link before saving','Enter a valid Drive sharing URL.'],['Anyone with the link can view','Change Drive sharing permission.'],['A date already exists','Edit the existing date record.'],['Upload failed','Retry; check type, size, connection, and authentication.'],['Saved but not public','Check Publish, then refresh the public page.'],['Old Google Sites link blocked','Replace with current site, Drive, YouTube/Vimeo, or a new upload.'],['Staff display broken','Validate the staff JSON array and save again.'],['Deleted file still exists on server','Expected behavior; media storage is immutable.']])

doc.add_heading('Recommended daily checklist',1)
bullets(['Choose the correct tab and record type before entering data.','Use consistent titles, dates, classes, subjects, and categories.','Test Drive links from a separate account before publishing.','Keep images under 15 MB and PDFs under 25 MB.','Use Draft while content is being reviewed.','After publishing, verify the exact public card, section, image, and link.','Use Delete only when the record should no longer exist; use Draft when it may be needed later.','Never place passwords or confidential student information in library fields.'])
doc.add_heading('Appendix: data and permissions',1)
tbl([['Area','Rule'],['librarySettings/main','Public profile configuration; administrators can create, update, or delete.'],['libraryPosts','Public reads only published records; authorised admins manage all records.'],['libraryOneFact','One date-based poster per document ID; public reads only published posters.'],['Media','Images use school server storage; PDFs use authenticated upload processing.'],['Library admin account','Must be authenticated and authorised.']])
para('End of manual').alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.core_properties.title='Library Admin Panel Instruction Manual'; doc.core_properties.subject='Ansar English School library administration'; doc.core_properties.author='Ansar English School'; doc.save(OUT); print(OUT)
